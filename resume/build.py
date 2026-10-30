"""Build the résumé from resume.yaml.

Usage (from the repo root):
  python resume/build.py              writes the .docx and public/resume.pdf
  python resume/build.py --docx-only  just the .docx

Needs: pip install python-docx pyyaml, and LibreOffice (soffice) on the PATH
for the PDF. The fonts are Carlito, which is metric-compatible with Calibri.
Layout matches the original LibreOffice résumé: Letter, 0.5" side margins,
10pt body, navy section headings with a rule underneath.
"""

import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: F401  (kept for future tables)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
DOCX = HERE / "Cameron-Carpenter-Resume.docx"
PDF = REPO / "public" / "resume.pdf"

FONT = "Carlito"
INK = RGBColor(0x1A, 0x1A, 0x1A)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x7A, 0x7A, 0x7A)
TEXT_WIDTH = Pt(539.9)          # 612pt page - 2 x 36pt margins
BULLET_TEXT = Pt(13.6)          # where wrapped bullet text lines up
BULLET_MARK = Pt(4.6)           # where the bullet itself sits

# Vertical rhythm, tuned against the original PDF.
LINE = Pt(12)
GAP_BULLET = Pt(0.4)
GAP_SECTION = Pt(4.4)


def run(p, text, size=10, bold=False, color=INK):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return r


def rich(p, text, size=10):
    """Text with **bold** spans."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            run(p, part, size=size, bold=(i % 2 == 1))


def link(p, text, url, size=9.5):
    rel = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    r = run(p, text, size=size)
    h.append(r._element)
    p._p.append(h)


def para(doc, align=None, before=0, after=0, line=LINE, exact=True):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_before, f.space_after = Pt(before) if not hasattr(before, "pt") else before, \
        Pt(after) if not hasattr(after, "pt") else after
    f.line_spacing = line
    f.line_spacing_rule = WD_LINE_SPACING.EXACTLY if exact else WD_LINE_SPACING.AT_LEAST
    f.widow_control = True
    if align is not None:
        p.alignment = align
    return p


def rule_below(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "1F3A5F"}.items():
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)


def heading(doc, text, first=False):
    p = para(doc, before=0 if first else GAP_SECTION, after=Pt(1.6), line=Pt(13.5))
    run(p, text.upper(), size=11, bold=True, color=NAVY)
    rule_below(p)
    p.paragraph_format.keep_with_next = True


def bullet(doc, text):
    p = para(doc, after=GAP_BULLET)
    f = p.paragraph_format
    f.left_indent, f.first_line_indent = BULLET_TEXT, -(BULLET_TEXT - BULLET_MARK)
    f.tab_stops.add_tab_stop(BULLET_TEXT)
    run(p, "•\t")
    rich(p, text)


def row(doc, left_bold, left_rest, right, before=Pt(3.0)):
    """'Org | Role' on the left, place and dates flush right."""
    p = para(doc, before=before, after=Pt(1))
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.keep_with_next = True
    run(p, left_bold, bold=True)
    if left_rest:
        run(p, left_rest)
    if right:
        run(p, "\t" + right)
    return p


def build(data):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(612), Pt(792)
    sec.left_margin = sec.right_margin = Pt(36)
    sec.top_margin, sec.bottom_margin = Pt(28.3), Pt(30)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10)
    doc.core_properties.title = f"{data['name']} Resume"
    doc.core_properties.author = data["name"]

    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, line=Pt(17.6), before=Pt(1.8))
    run(p, data["name"], size=20, bold=True, color=NAVY)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, line=Pt(12.5))
    run(p, data["tagline"], size=11)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, line=Pt(13), after=Pt(4.1))
    for i, c in enumerate(data["contact"]):
        if i:
            run(p, "  |  ", size=9.5, color=GREY)
        if c.get("link"):
            link(p, c["text"], c["link"])
        else:
            run(p, c["text"], size=9.5)

    heading(doc, "Summary", first=True)
    p = para(doc)
    rich(p, data["summary"])

    heading(doc, "Experience")
    for i, job in enumerate(data["experience"]):
        row(doc, f"{job['org']} | {job['role']}", "", f"{job['where']} | {job['when']}",
            before=Pt(0.8) if i == 0 else Pt(3.0))
        for b in job["bullets"]:
            bullet(doc, b)

    heading(doc, "Projects")
    for i, proj in enumerate(data["projects"]):
        row(doc, f"{proj['name']} | {proj['note']}", "", "", before=Pt(0.8) if i == 0 else Pt(3.0))
        for b in proj["bullets"]:
            bullet(doc, b)

    heading(doc, "Education")
    for e in data["education"]:
        p = row(doc, e["school"], f" | {e['degree']}", e["when"], before=Pt(0.8))

    heading(doc, "Skills")
    for s in data["skills"]:
        bullet(doc, s)
    return doc


def main():
    data = yaml.safe_load(open(HERE / "resume.yaml", encoding="utf-8"))
    doc = build(data)
    doc.save(DOCX)
    print(f"Wrote {DOCX.relative_to(REPO)}")
    if "--docx-only" in sys.argv:
        return
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("LibreOffice (soffice) not found, so no PDF. Open the .docx and export to PDF instead.")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(DOCX)],
                       check=True, capture_output=True)
        shutil.copy(Path(tmp) / (DOCX.stem + ".pdf"), PDF)
    print(f"Wrote {PDF.relative_to(REPO)}")


if __name__ == "__main__":
    main()
